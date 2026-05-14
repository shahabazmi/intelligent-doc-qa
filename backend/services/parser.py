import os
from pathlib import Path


def parse_document(file_path: str) -> list[dict]:
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            return _parse_pdf(file_path)
        elif ext == ".docx":
            return _parse_docx(file_path)
        elif ext == ".txt":
            return _parse_txt(file_path)
        elif ext in {".png", ".jpg", ".jpeg"}:
            return _parse_image(file_path)
    except Exception as e:
        raise ValueError(f"Failed to parse {os.path.basename(file_path)}: {e}") from e
    raise ValueError(f"Unsupported file type: {ext}")


def _parse_pdf(file_path: str) -> list[dict]:
    import fitz  # pymupdf
    doc = fitz.open(file_path)
    parsed = []
    filename = os.path.basename(file_path)
    for page_num, page in enumerate(doc, start=1):
        # Extract tables first (PyMuPDF 1.23+) — preserves columnar structure for
        # annual reports, financial statements, structured PDFs.
        try:
            finder = page.find_tables()
            tables = getattr(finder, "tables", []) or []
        except Exception:
            tables = []
        for table_idx, table in enumerate(tables, start=1):
            try:
                rows = table.extract() or []
            except Exception:
                continue
            cleaned = [
                [("" if c is None else str(c)).strip() for c in row]
                for row in rows
                if row and any((c or "") for c in row)
            ]
            if not cleaned:
                continue
            ncols = max(len(r) for r in cleaned)
            cleaned = [r + [""] * (ncols - len(r)) for r in cleaned]
            header = " | ".join(cleaned[0])
            sep    = " | ".join(["---"] * ncols)
            body   = "\n".join(" | ".join(r) for r in cleaned[1:])
            table_text = (
                f"[TABLE {table_idx} — page {page_num}]\n"
                f"{header}\n{sep}\n{body}\n"
                f"[/TABLE]"
            ).strip()
            parsed.append({
                "text": table_text,
                "metadata": {"category": "Table", "page_number": page_num, "filename": filename},
            })

        text = page.get_text().strip()
        if text:
            parsed.append({
                "text": text,
                "metadata": {"category": "Text", "page_number": page_num, "filename": filename},
            })
    doc.close()
    return parsed


def _parse_docx(file_path: str) -> list[dict]:
    from docx import Document
    doc = Document(file_path)
    filename = os.path.basename(file_path)
    parsed = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parsed.append({
                "text": text,
                "metadata": {"category": "Text", "page_number": 1, "filename": filename},
            })
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parsed.append({
                    "text": " | ".join(cells),
                    "metadata": {"category": "Table", "page_number": 1, "filename": filename},
                })
    return parsed


def _parse_txt(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    with open(file_path, encoding="utf-8", errors="replace") as f:
        text = f.read().strip()
    if not text:
        return []
    return [{"text": text, "metadata": {"category": "Text", "page_number": 1, "filename": filename}}]


def _parse_image(file_path: str) -> list[dict]:
    try:
        from unstructured.partition.auto import partition
        elements = partition(filename=file_path, strategy="auto")
        filename = os.path.basename(file_path)
        parsed = []
        for el in elements:
            text = str(el).strip()
            if text:
                metadata = getattr(el, "metadata", None)
                page = (getattr(metadata, "page_number", None) or 1) if metadata else 1
                parsed.append({
                    "text": text,
                    "metadata": {"category": str(getattr(el, "category", "")), "page_number": page, "filename": filename},
                })
        return parsed
    except Exception as e:
        raise ValueError(f"Image parsing failed: {e}") from e
